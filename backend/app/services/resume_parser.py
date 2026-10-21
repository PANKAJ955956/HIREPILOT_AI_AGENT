import re
from pathlib import Path
import fitz
from docx import Document

SKILL_TERMS = ["python", "sql", "machine learning", "deep learning", "generative ai", "agentic ai", "langchain", "langgraph", "rag", "retrieval augmented generation", "fastapi", "flask", "django", "pandas", "numpy", "scikit-learn", "pytorch", "tensorflow", "transformers", "llm", "openai", "gemini", "chromadb", "faiss", "postgresql", "mysql", "mongodb", "docker", "aws", "azure", "gcp", "react", "typescript", "javascript", "power bi", "tableau", "excel", "data analysis", "nlp", "computer vision", "rest api", "git", "linux", "statistics", "data structures", "algorithms"]

def extract_text(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        with fitz.open(stream=content, filetype="pdf") as doc:
            text = "\n".join(page.get_text("text") for page in doc)
    elif suffix == ".docx":
        import io
        doc = Document(io.BytesIO(content))
        text = "\n".join(p.text for p in doc.paragraphs)
        for table in doc.tables:
            text += "\n" + "\n".join(" | ".join(c.text for c in row.cells) for row in table.rows)
    elif suffix == ".txt":
        text = content.decode("utf-8", errors="replace")
    else:
        raise ValueError("Unsupported file type. Upload PDF, DOCX, or TXT.")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) < 40:
        raise ValueError("Not enough extractable text. If this is a scanned PDF, OCR it first and upload a text-based PDF.")
    return text[:100_000]

def extract_skills(text: str) -> list[str]:
    normalized = text.lower()
    found = []
    for skill in SKILL_TERMS:
        if re.search(r"(?<![a-z0-9])" + re.escape(skill) + r"(?![a-z0-9])", normalized):
            found.append(skill)
    return sorted(set(found))

def extract_profile(text: str, filename: str) -> dict:
    email_match = re.search(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", text, re.I)
    lines = [x.strip() for x in text.splitlines() if x.strip()]
    name = lines[0][:200] if lines and len(lines[0]) < 100 and "@" not in lines[0] else "Unknown candidate"
    education = [line[:300] for line in lines if re.search(r"b\.?tech|bachelor|master|bsc|m\.?tech|university|college|degree|education|ph\.?d", line, re.I)][:12]
    experience = [line for line in lines if re.search(r"experience|intern|engineer|developer|analyst|manager|worked|employment", line, re.I)][:12]
    return {"full_name": name, "email": email_match.group(0) if email_match else None, "skills": extract_skills(text), "education": education, "experience_summary": "\n".join(experience)[:4000], "resume_filename": filename, "resume_text": text}
